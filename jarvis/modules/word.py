"""MS Word automation — python-docx file engine + Word.Application COM app engine.

File engine (cross-platform): create, open, read, headings/body/tables, find &
replace, save — real .docx artifacts via the Open XML format.
App engine (Windows + MS Word): launch the document in the real Word UI and
export to PDF through Word.Application. Shadows to dry-run logs elsewhere.
"""

from __future__ import annotations

from pathlib import Path

from ..exceptions import ActionError
from ..logging_setup import get_logger
from .office import OfficeModule

log = get_logger("jarvis.word")


class WordModule(OfficeModule):
    name = "word"
    progid = "Word.Application"
    file_module = "docx"
    app_label = "MS Word"

    # ── engine plumbing ──────────────────────────────────────────────────
    def _docx(self):
        self.require_real()
        import docx

        return docx

    def _ensure_doc(self):
        if self._doc is None:
            self._doc = self._docx().Document()
        return self._doc

    # ── file engine: lifecycle ───────────────────────────────────────────
    def create(self, path: str | None = None, title: str | None = None) -> dict:
        if self.dry_run:
            return self.shadow("create", path=path, title=title)
        self._doc = self._docx().Document()
        if title:
            self._doc.core_properties.title = title
        self._path = Path(path) if path else None
        return {"created": str(self._path) if self._path else "(unsaved)",
                "title": title}

    def open(self, path: str) -> dict:
        if self.dry_run:
            return self.shadow("open", path=path)
        p = Path(path).expanduser()
        if not p.exists():
            raise ActionError(f"word: document not found: {p}")
        self._doc = self._docx().Document(str(p))
        self._path = p
        return {"opened": str(p), "paragraphs": len(self._doc.paragraphs),
                "tables": len(self._doc.tables)}

    def save(self, path: str | None = None) -> dict:
        if self.dry_run:
            return self.shadow("save", path=path)
        if self._doc is None:
            raise ActionError("word.save: no document open or created")
        p = self._target_path(path, suffix=".docx")
        p.parent.mkdir(parents=True, exist_ok=True)
        self._doc.save(str(p))
        self._path = p
        log.info("saved %s (%d bytes)", p, p.stat().st_size)
        return {"saved": str(p), "bytes": p.stat().st_size}

    # ── file engine: inspection ──────────────────────────────────────────
    def read(self) -> dict:
        """Structure inventory: paragraphs, headings, tables, word count."""
        if self.dry_run:
            return self.shadow("read")
        doc = self._ensure_doc()
        paragraphs, headings, words = [], [], 0
        for p in doc.paragraphs:
            style = p.style.name if p.style else ""
            paragraphs.append({"text": p.text, "style": style})
            if style.startswith("Heading"):
                try:
                    level = int(style.split()[-1])
                except ValueError:  # pragma: no cover - non-numeric heading styles
                    level = 1
                headings.append({"level": level, "text": p.text})
            words += len(p.text.split())
        tables = [[[cell.text for cell in row.cells] for row in t.rows]
                  for t in doc.tables]
        for table in tables:
            for row in table:
                words += sum(len(str(cell).split()) for cell in row)
        return {
            "path": str(self._path) if self._path else None,
            "title": doc.core_properties.title or "",
            "paragraphs": paragraphs,
            "headings": headings,
            "tables": tables,
            "word_count": words,
        }

    # ── file engine: editing ─────────────────────────────────────────────
    def add_heading(self, text: str, level: int = 1) -> dict:
        if self.dry_run:
            return self.shadow("add_heading", text=text, level=level)
        level = max(1, min(int(level), 4))
        self._ensure_doc().add_heading(text, level=level)
        return {"added_heading": text, "level": level}

    def add_text(self, text: str, style: str | None = None,
                 bold: bool = False) -> dict:
        if self.dry_run:
            return self.shadow("add_text", text=text, style=style, bold=bold)
        doc = self._ensure_doc()
        p = doc.add_paragraph()
        run = p.add_run(text)
        run.bold = bool(bold)
        if style:
            try:
                p.style = style
            except (KeyError, ValueError) as exc:
                raise ActionError(f"word.add_text: unknown style {style!r}") from exc
        return {"added_text_chars": len(text), "style": style, "bold": bool(bold)}

    def add_table(self, rows: list[list]) -> dict:
        if self.dry_run:
            return self.shadow("add_table", rows=rows)
        if not rows or not isinstance(rows, list):
            raise ActionError("word.add_table: rows must be a non-empty list of rows")
        grid = [[str(c) for c in row] for row in rows]
        width = max(len(r) for r in grid)
        table = self._ensure_doc().add_table(rows=len(grid), cols=width)
        try:
            table.style = "Table Grid"
        except (KeyError, ValueError):  # pragma: no cover - style always in default tpl
            pass
        for i, row in enumerate(grid):
            for j, value in enumerate(row):
                table.rows[i].cells[j].text = value
        return {"added_table": f"{len(grid)}x{width}"}

    def replace(self, find: str, replace: str = "") -> dict:
        """Find & replace across body paragraphs and table cells (run-local)."""
        if self.dry_run:
            return self.shadow("replace", find=find, replace=replace)
        if not find:
            raise ActionError("word.replace: find string must be non-empty")
        doc = self._ensure_doc()
        count = 0

        def swap(paragraphs) -> None:
            nonlocal count
            for p in paragraphs:
                for run in p.runs:
                    if find in run.text:
                        count += run.text.count(find)
                        run.text = run.text.replace(find, replace)

        swap(doc.paragraphs)
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    swap(cell.paragraphs)
        return {"replaced": count, "find": find, "replace": replace}

    # ── app engine: real MS Word (COM) ───────────────────────────────────
    def export_pdf(self, path: str | None = None) -> dict:
        """Export the document to PDF through real MS Word."""
        target = self._target_path(path).with_suffix(".pdf")
        sh = self.app_shadow("export_pdf", path=str(target))
        if sh is not None:
            return sh
        if self._path is None:
            raise ActionError("word.export_pdf: save the document first")
        word = self._app()
        word.Visible = False
        doc = word.Documents.Open(self._abs(self._path))
        try:
            doc.SaveAs2(self._abs(target), FileFormat=17)  # 17 = wdFormatPDF
        finally:
            doc.Close(False)
        return {"exported_pdf": str(target), "bytes": target.stat().st_size}

    def launch(self, path: str | None = None) -> dict:
        """Open the document in the real MS Word application window."""
        p = self._target_path(path, suffix=".docx")
        sh = self.app_shadow("launch", path=str(p))
        if sh is not None:
            return sh
        word = self._app()
        word.Visible = True
        word.Documents.Open(self._abs(p))
        return {"launched": str(p), "app": self.progid}
