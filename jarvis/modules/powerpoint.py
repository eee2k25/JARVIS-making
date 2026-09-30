"""MS PowerPoint automation — python-pptx file engine + PowerPoint COM app engine.

File engine (cross-platform): create, open, read slide inventories, add
title/content slides with speaker notes, embed images, save — real .pptx
artifacts via the Open XML format.
App engine (Windows + MS PowerPoint): present/open the deck in the real
PowerPoint UI and export to PDF through PowerPoint.Application. Shadows to
dry-run logs elsewhere.
"""

from __future__ import annotations

from pathlib import Path

from ..exceptions import ActionError
from ..logging_setup import get_logger
from .office import OfficeModule

log = get_logger("jarvis.powerpoint")


class PowerPointModule(OfficeModule):
    name = "powerpoint"
    progid = "PowerPoint.Application"
    file_module = "pptx"
    app_label = "MS PowerPoint"

    # ── engine plumbing ──────────────────────────────────────────────────
    def _pptx(self):
        self.require_real()
        import pptx

        return pptx

    def _ensure_deck(self):
        if self._doc is None:
            self._doc = self._pptx().Presentation()
        return self._doc

    # ── file engine: lifecycle ───────────────────────────────────────────
    def create(self, path: str | None = None, title: str | None = None) -> dict:
        if self.dry_run:
            return self.shadow("create", path=path, title=title)
        self._doc = self._pptx().Presentation()
        if title:
            self._doc.core_properties.title = title
        self._path = Path(path) if path else None
        return {"created": str(self._path) if self._path else "(unsaved)",
                "title": title, "slides": 0}

    def open(self, path: str) -> dict:
        if self.dry_run:
            return self.shadow("open", path=path)
        p = Path(path).expanduser()
        if not p.exists():
            raise ActionError(f"powerpoint: deck not found: {p}")
        self._doc = self._pptx().Presentation(str(p))
        self._path = p
        return {"opened": str(p), "slides": len(self._doc.slides)}

    def save(self, path: str | None = None) -> dict:
        if self.dry_run:
            return self.shadow("save", path=path)
        if self._doc is None:
            raise ActionError("powerpoint.save: no presentation open or created")
        p = self._target_path(path, suffix=".pptx")
        p.parent.mkdir(parents=True, exist_ok=True)
        self._doc.save(str(p))
        self._path = p
        log.info("saved %s (%d bytes)", p, p.stat().st_size)
        return {"saved": str(p), "bytes": p.stat().st_size}

    # ── file engine: inspection ──────────────────────────────────────────
    def read(self) -> dict:
        """Slide inventory: titles, bullets, notes per slide."""
        if self.dry_run:
            return self.shadow("read")
        deck = self._ensure_deck()
        slides = []
        for idx, slide in enumerate(deck.slides, start=1):
            title_shape = slide.shapes.title
            title = title_shape.text if title_shape else ""
            bullets: list[str] = []
            for shape in slide.placeholders:
                if title_shape is not None and shape.shape_id == title_shape.shape_id:
                    continue
                if not shape.has_text_frame:
                    continue
                for para in shape.text_frame.paragraphs:
                    text = para.text.strip()
                    if text:
                        bullets.append(text)
            notes = ""
            if slide.has_notes_slide:
                notes = slide.notes_slide.notes_text_frame.text.strip()
            slides.append({"index": idx, "title": title,
                           "bullets": bullets, "notes": notes})
        return {
            "path": str(self._path) if self._path else None,
            "title": deck.core_properties.title or "",
            "slide_count": len(slides),
            "slides": slides,
        }

    # ── file engine: editing ─────────────────────────────────────────────
    def add_slide(self, title: str, bullets: list[str] | None = None,
                  notes: str | None = None) -> dict:
        if self.dry_run:
            return self.shadow("add_slide", title=title, bullets=bullets,
                               notes=notes)
        deck = self._ensure_deck()
        layout = deck.slide_layouts[1]  # "Title and Content"
        slide = deck.slides.add_slide(layout)
        if title:
            slide.shapes.title.text = title
        bullets = [str(b) for b in (bullets or []) if str(b).strip()]
        if bullets:
            body = slide.placeholders[1].text_frame
            body.text = bullets[0]
            for bullet in bullets[1:]:
                p = body.add_paragraph()
                p.text = bullet
        if notes:
            slide.notes_slide.notes_text_frame.text = notes
        return {"slide": len(deck.slides), "title": title,
                "bullets": len(bullets), "notes": bool(notes)}

    def add_image(self, path: str, width_in: float = 8.0) -> dict:
        if self.dry_run:
            return self.shadow("add_image", path=path, width_in=width_in)
        img = Path(path).expanduser()
        if not img.exists():
            raise ActionError(f"powerpoint.add_image: image not found: {img}")
        deck = self._ensure_deck()
        if not deck.slides:
            raise ActionError("powerpoint.add_image: add a slide first")
        from pptx.util import Inches

        picture = deck.slides[-1].shapes.add_picture(
            str(img), Inches(1), Inches(2), width=Inches(float(width_in)))
        return {"added_image": str(img), "slide": len(deck.slides),
                "width_in": float(width_in), "shape_id": picture.shape_id}

    # ── app engine: real MS PowerPoint (COM) ─────────────────────────────
    def export_pdf(self, path: str | None = None) -> dict:
        """Export the deck to PDF through real MS PowerPoint."""
        target = self._target_path(path).with_suffix(".pdf")
        sh = self.app_shadow("export_pdf", path=str(target))
        if sh is not None:
            return sh
        if self._path is None:
            raise ActionError("powerpoint.export_pdf: save the deck first")
        app = self._app()
        deck = app.Presentations.Open(self._abs(self._path), WithWindow=False)
        try:
            deck.SaveAs(self._abs(target), 32)  # 32 = ppSaveAsPDF
        finally:
            deck.Close()
        return {"exported_pdf": str(target), "bytes": target.stat().st_size}

    def launch(self, path: str | None = None) -> dict:
        """Open the deck in the real MS PowerPoint application window."""
        p = self._target_path(path, suffix=".pptx")
        sh = self.app_shadow("launch", path=str(p))
        if sh is not None:
            return sh
        app = self._app()
        app.Visible = True
        app.Presentations.Open(self._abs(p), WithWindow=True)
        return {"launched": str(p), "app": self.progid}
